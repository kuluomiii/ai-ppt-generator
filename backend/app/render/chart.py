import logging

from pptx.chart.chart import Chart
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.slide import Slide as PptxSlide
from pptx.util import Pt

from app.domain.content import ChartBlock, ChartKind
from app.domain.geometry import Rect
from app.domain.theme import TextStyle, Theme
from app.render.color import to_rgb
from app.render.text import apply_font_style, apply_text_style

logger = logging.getLogger(__name__)

_CHART_TYPES: dict[ChartKind, XL_CHART_TYPE] = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}

# python-pptx 不暴露图表区/绘图区/图例的填充与边框 API，只能写 spPr。
_TRANSPARENT_SPPR = (
    '<c:spPr xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    "<a:noFill/><a:ln><a:noFill/></a:ln>"
    "</c:spPr>"
)

_GRIDLINE_WIDTH_PT = 0.75
_LINE_SERIES_WIDTH_PT = 2.25


def render_chart(
    pptx_slide: PptxSlide,
    rect: Rect,
    block: ChartBlock,
    theme: Theme,
) -> None:
    categories, series_rows = _align_chart_data(block)
    if not categories or not series_rows:
        logger.warning(
            "图表数据为空，跳过渲染：categories=%s series=%s",
            len(block.categories),
            len(block.series),
        )
        return

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series_rows:
        chart_data.add_series(name, values)

    left, top, width, height = (Pt(value) for value in rect.to_points())
    graphic_frame = pptx_slide.shapes.add_chart(
        _CHART_TYPES[block.chart_type],
        left,
        top,
        width,
        height,
        chart_data,
    )
    chart = graphic_frame.chart
    _theme_chart(chart, block, theme, series_count=len(series_rows))


def _align_chart_data(block: ChartBlock) -> tuple[list[str], list[tuple[str, tuple[float, ...]]]]:
    if not block.series:
        logger.warning("图表 series 为空，无法渲染")
        return [], []

    lengths = [len(block.categories), *[len(item.values) for item in block.series]]
    count = min(lengths)
    if count != len(block.categories) or any(
        len(item.values) != len(block.categories) for item in block.series
    ):
        logger.warning(
            "图表分类与系列长度不一致，按较短长度对齐：categories=%s series_lengths=%s",
            len(block.categories),
            [len(item.values) for item in block.series],
        )

    if count == 0:
        return [], []

    categories = block.categories[:count]
    series_rows = [(item.name, tuple(item.values[:count])) for item in block.series]
    return categories, series_rows


def _theme_chart(
    chart: Chart,
    block: ChartBlock,
    theme: Theme,
    *,
    series_count: int,
) -> None:
    label_style = theme.text_style("chart_label")
    colors = theme.palette.chart_series

    chart.has_title = False
    apply_font_style(chart.font, theme, label_style)
    _clear_chart_chrome(chart)

    if block.chart_type == "pie":
        _style_pie(chart, theme, colors, label_style)
    else:
        _style_cartesian(chart, block, theme, colors, label_style)

    if series_count > 1:
        chart.has_legend = True
        chart.legend.include_in_layout = False
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        apply_font_style(chart.legend.font, theme, label_style)
        _set_transparent_sppr(chart.legend._element)
    else:
        chart.has_legend = False


def _style_cartesian(
    chart: Chart,
    block: ChartBlock,
    theme: Theme,
    colors: list[str],
    label_style: TextStyle,
) -> None:
    for index, series in enumerate(chart.series):
        _color_series(series, colors[index % len(colors)], block.chart_type)

    plot = chart.plots[0]
    plot.has_data_labels = False

    category_axis = chart.category_axis
    category_axis.has_major_gridlines = False
    category_axis.has_minor_gridlines = False
    category_axis.major_tick_mark = XL_TICK_MARK.NONE
    category_axis.minor_tick_mark = XL_TICK_MARK.NONE
    category_axis.format.line.fill.background()
    apply_font_style(category_axis.tick_labels.font, theme, label_style)

    value_axis = chart.value_axis
    value_axis.has_major_gridlines = True
    value_axis.has_minor_gridlines = False
    value_axis.major_tick_mark = XL_TICK_MARK.NONE
    value_axis.minor_tick_mark = XL_TICK_MARK.NONE
    value_axis.format.line.fill.background()
    apply_font_style(value_axis.tick_labels.font, theme, label_style)

    gridline = value_axis.major_gridlines
    gridline.format.line.color.rgb = to_rgb(theme.palette.line)
    gridline.format.line.width = Pt(_GRIDLINE_WIDTH_PT)

    if block.unit:
        value_axis.has_title = True
        paragraph = value_axis.axis_title.text_frame.paragraphs[0]
        paragraph.clear()
        run = paragraph.add_run()
        run.text = block.unit
        apply_text_style(run, theme, label_style)
    else:
        value_axis.has_title = False


def _style_pie(
    chart: Chart,
    theme: Theme,
    colors: list[str],
    label_style: TextStyle,
) -> None:
    series = chart.series[0]
    for index, point in enumerate(series.points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = to_rgb(colors[index % len(colors)])

    plot = chart.plots[0]
    plot.has_data_labels = True
    data_labels = plot.data_labels
    data_labels.show_percentage = True
    data_labels.show_value = False
    data_labels.show_category_name = False
    data_labels.show_series_name = False
    apply_font_style(data_labels.font, theme, label_style)


def _color_series(series, color: str, chart_type: ChartKind) -> None:
    rgb = to_rgb(color)
    if chart_type == "line":
        # 折线的可见颜色在线条上；填充只影响标记，两者都设以免主题色丢失
        series.format.line.color.rgb = rgb
        series.format.line.width = Pt(_LINE_SERIES_WIDTH_PT)
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = rgb
        series.marker.format.line.color.rgb = rgb
        return

    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb


def _clear_chart_chrome(chart: Chart) -> None:
    chart_space = chart._element
    chart_el = chart_space.find(qn("c:chart"))
    if chart_el is None:
        return

    _set_transparent_sppr(chart_space, before=qn("c:txPr"))
    plot_area = chart_el.find(qn("c:plotArea"))
    if plot_area is not None:
        # CT_PlotArea: layout?, (chart group)+, (axis)*, dTable?, spPr?
        _set_transparent_sppr(plot_area)


def _set_transparent_sppr(parent, *, before: str | None = None) -> None:
    for existing in parent.findall(qn("c:spPr")):
        parent.remove(existing)

    sp_pr = parse_xml(_TRANSPARENT_SPPR)
    if before is not None:
        anchor = parent.find(before)
        if anchor is not None:
            anchor.addprevious(sp_pr)
            return
    parent.append(sp_pr)
