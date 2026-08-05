"""PPTX 回读与可编辑性验证。

输入渲染结果字节流与源 Deck，检查页数、画布、原生对象与内容完整性。
供导出接口与第 14 节点回归脚本复用，不依赖 HTTP 层。
"""

from __future__ import annotations

from io import BytesIO
from typing import BinaryIO

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.presentation import Presentation as PresentationType
from pptx.shapes.base import BaseShape
from pptx.slide import Slide as PptxSlide
from pptx.util import Emu, Pt
from pydantic import BaseModel, Field

from app.domain.content import (
    Block,
    BulletsBlock,
    ChartBlock,
    Deck,
    ImageBlock,
    KpiBlock,
    Slide,
    TableBlock,
    TextBlock,
)
from app.domain.geometry import CANVAS_HEIGHT_PT, CANVAS_WIDTH_PT

# 半 pt 容差：渲染取整与自由形状边界可能有亚点误差
_BOUNDS_TOLERANCE_EMU = Emu(Pt(0.5))
# 覆盖整页的判据：宽高与面积均达到画布的 90% 以上
_FULL_PAGE_RATIO = 0.9
_CANVAS_TOLERANCE_EMU = Emu(Pt(0.5))


class VerifyIssue(BaseModel):
    """单条可编辑性/完整性问题，可定位到页与形状。"""

    check: str
    slide_index: int | None = None
    shape: str | None = None
    message: str


class VerifyReport(BaseModel):
    """回读验证报告。"""

    passed: bool
    issues: list[VerifyIssue] = Field(default_factory=list)
    slide_count: int = 0
    expected_slide_count: int = 0


def verify_pptx(data: bytes | BinaryIO, deck: Deck) -> VerifyReport:
    """用 python-pptx 重新打开导出结果并核对源 Deck。"""
    stream = BytesIO(data) if isinstance(data, (bytes, bytearray)) else data
    if hasattr(stream, "seek"):
        stream.seek(0)

    issues: list[VerifyIssue] = []
    try:
        presentation = Presentation(stream)
    except Exception as error:
        return VerifyReport(
            passed=False,
            issues=[
                VerifyIssue(
                    check="open",
                    message=f"无法用 python-pptx 打开导出文件：{error}",
                )
            ],
            expected_slide_count=len(deck.slides),
        )

    expected = len(deck.slides)
    actual = len(presentation.slides)
    issues.extend(_check_page_count(actual, expected))
    issues.extend(_check_canvas_size(presentation))

    # 页数不一致时仍尽量检查已有页，避免一次失败掩盖其它问题
    for index, (source, pptx_slide) in enumerate(
        zip(deck.slides, presentation.slides, strict=False),
        start=1,
    ):
        issues.extend(_check_slide(index, source, pptx_slide, presentation))

    return VerifyReport(
        passed=not issues,
        issues=issues,
        slide_count=actual,
        expected_slide_count=expected,
    )


def _check_page_count(actual: int, expected: int) -> list[VerifyIssue]:
    if actual == expected:
        return []
    return [
        VerifyIssue(
            check="page_count",
            message=f"页数不一致：导出 {actual} 页，项目应为 {expected} 页",
        )
    ]


def _check_canvas_size(presentation: PresentationType) -> list[VerifyIssue]:
    expected_w = Emu(Pt(CANVAS_WIDTH_PT))
    expected_h = Emu(Pt(CANVAS_HEIGHT_PT))
    width = presentation.slide_width
    height = presentation.slide_height
    if width is None or height is None:
        return [
            VerifyIssue(
                check="canvas_size",
                message="页面尺寸缺失，无法确认是否为 16:9",
            )
        ]

    if abs(int(width) - int(expected_w)) > int(_CANVAS_TOLERANCE_EMU) or abs(
        int(height) - int(expected_h)
    ) > int(_CANVAS_TOLERANCE_EMU):
        ratio = (int(width) / int(height)) if int(height) else 0.0
        return [
            VerifyIssue(
                check="canvas_size",
                message=(
                    f"页面尺寸不是约定的 16:9（960×540 pt）："
                    f"当前约 {int(width) / 12700:.1f}×{int(height) / 12700:.1f} pt，"
                    f"宽高比 {ratio:.3f}"
                ),
            )
        ]
    return []


def _check_slide(
    slide_index: int,
    source: Slide,
    pptx_slide: PptxSlide,
    presentation: PresentationType,
) -> list[VerifyIssue]:
    issues: list[VerifyIssue] = []
    shapes = list(pptx_slide.shapes)

    issues.extend(_check_full_page_pictures(slide_index, shapes, presentation))
    issues.extend(_check_shape_bounds(slide_index, shapes, presentation))
    issues.extend(_check_native_tables(slide_index, source, shapes))
    issues.extend(_check_native_charts(slide_index, source, shapes))
    issues.extend(_check_text_in_frames(slide_index, source, shapes))
    issues.extend(_check_content_integrity(slide_index, source, shapes))
    return issues


def _check_full_page_pictures(
    slide_index: int,
    shapes: list[BaseShape],
    presentation: PresentationType,
) -> list[VerifyIssue]:
    slide_w = int(presentation.slide_width or 0)
    slide_h = int(presentation.slide_height or 0)
    if slide_w <= 0 or slide_h <= 0:
        return []

    slide_area = slide_w * slide_h
    issues: list[VerifyIssue] = []
    for shape in shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        width = int(shape.width)
        height = int(shape.height)
        area = width * height
        if (
            width >= slide_w * _FULL_PAGE_RATIO
            and height >= slide_h * _FULL_PAGE_RATIO
            and area >= slide_area * _FULL_PAGE_RATIO
        ):
            issues.append(
                VerifyIssue(
                    check="full_page_picture",
                    slide_index=slide_index,
                    shape=_shape_label(shape),
                    message=(
                        f"第 {slide_index} 页存在覆盖整页的图片"
                        f"（{_shape_label(shape)}），疑似用截图承载正文"
                    ),
                )
            )
    return issues


def _check_shape_bounds(
    slide_index: int,
    shapes: list[BaseShape],
    presentation: PresentationType,
) -> list[VerifyIssue]:
    slide_w = int(presentation.slide_width or 0)
    slide_h = int(presentation.slide_height or 0)
    tol = int(_BOUNDS_TOLERANCE_EMU)
    issues: list[VerifyIssue] = []

    for shape in shapes:
        left = int(shape.left)
        top = int(shape.top)
        right = left + int(shape.width)
        bottom = top + int(shape.height)
        if left < -tol or top < -tol or right > slide_w + tol or bottom > slide_h + tol:
            issues.append(
                VerifyIssue(
                    check="bounds",
                    slide_index=slide_index,
                    shape=_shape_label(shape),
                    message=(
                        f"第 {slide_index} 页形状越界：{_shape_label(shape)}"
                        f"（left={left}, top={top}, right={right}, bottom={bottom}；"
                        f"画布 {slide_w}×{slide_h} EMU）"
                    ),
                )
            )
    return issues


def _check_native_tables(
    slide_index: int,
    source: Slide,
    shapes: list[BaseShape],
) -> list[VerifyIssue]:
    expected = sum(1 for block in source.blocks if isinstance(block, TableBlock))
    if expected == 0:
        return []

    actual = sum(1 for shape in shapes if shape.has_table)
    if actual >= expected:
        return []
    return [
        VerifyIssue(
            check="native_table",
            slide_index=slide_index,
            message=(
                f"第 {slide_index} 页缺少原生表格：源内容有 {expected} 个表格块，"
                f"导出仅找到 {actual} 个 GraphicFrame 表格"
            ),
        )
    ]


def _check_native_charts(
    slide_index: int,
    source: Slide,
    shapes: list[BaseShape],
) -> list[VerifyIssue]:
    expected = sum(1 for block in source.blocks if isinstance(block, ChartBlock))
    if expected == 0:
        return []

    actual = sum(1 for shape in shapes if shape.has_chart)
    if actual >= expected:
        return []
    return [
        VerifyIssue(
            check="native_chart",
            slide_index=slide_index,
            message=(
                f"第 {slide_index} 页缺少原生图表：源内容有 {expected} 个图表块，"
                f"导出仅找到 {actual} 个 GraphicFrame 图表"
            ),
        )
    ]


def _check_text_in_frames(
    slide_index: int,
    source: Slide,
    shapes: list[BaseShape],
) -> list[VerifyIssue]:
    """正文必须落在真正的文本框里，而不是位图。"""
    expected_texts = [
        text
        for block in source.blocks
        if isinstance(block, (TextBlock, BulletsBlock, KpiBlock))
        for text in key_texts_from_block(block)
    ]
    if not expected_texts:
        return []

    frame_texts = _collect_text_frame_strings(shapes)
    if not frame_texts:
        return [
            VerifyIssue(
                check="text_frame",
                slide_index=slide_index,
                message=(
                    f"第 {slide_index} 页有文字内容，但未找到带非空文字的文本框（可能被位图替代）"
                ),
            )
        ]

    issues: list[VerifyIssue] = []
    for text in expected_texts:
        if not _text_present(text, frame_texts):
            issues.append(
                VerifyIssue(
                    check="text_frame",
                    slide_index=slide_index,
                    message=(f"第 {slide_index} 页文字未出现在文本框中：「{_preview(text)}」"),
                )
            )
    return issues


def _check_content_integrity(
    slide_index: int,
    source: Slide,
    shapes: list[BaseShape],
) -> list[VerifyIssue]:
    haystack = _collect_all_strings(shapes)
    issues: list[VerifyIssue] = []
    for block in source.blocks:
        for text in key_texts_from_block(block):
            if _text_present(text, haystack):
                continue
            issues.append(
                VerifyIssue(
                    check="content_integrity",
                    slide_index=slide_index,
                    message=(
                        f"第 {slide_index} 页内容缺失：块 {block.id}"
                        f"（{block.type}）的关键文字「{_preview(text)}」未在导出结果中找到"
                    ),
                )
            )
    return issues


def key_texts_from_block(block: Block) -> list[str]:
    """块级关键文字，供完整性核对与回归脚本复用。"""
    match block:
        case TextBlock(text=text):
            return [text] if text.strip() else []
        case BulletsBlock(items=items):
            return [item for item in items if item.strip()]
        case KpiBlock(value=value, label=label, note=note):
            texts = [value, label]
            if note:
                texts.append(note)
            return [item for item in texts if item.strip()]
        case TableBlock(header=header, rows=rows):
            texts = list(header)
            for row in rows:
                texts.extend(row)
            return [item for item in texts if item.strip()]
        case ChartBlock(categories=categories, series=series):
            texts = list(categories)
            texts.extend(item.name for item in series if item.name)
            return [item for item in texts if item.strip()]
        case ImageBlock():
            return []
        case _:
            return []


def _collect_text_frame_strings(shapes: list[BaseShape]) -> list[str]:
    texts: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            continue
        if shape.has_table or shape.has_chart:
            continue
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            value = paragraph.text.strip()
            if value:
                texts.append(value)
    return texts


def _collect_all_strings(shapes: list[BaseShape]) -> list[str]:
    texts = list(_collect_text_frame_strings(shapes))
    for shape in shapes:
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    value = cell.text.strip()
                    if value:
                        texts.append(value)
        if shape.has_chart:
            texts.extend(_chart_strings(shape))
    return texts


def _chart_strings(shape: BaseShape) -> list[str]:
    chart = shape.chart
    texts: list[str] = []
    try:
        for plot in chart.plots:
            for category in plot.categories:
                value = str(category).strip()
                if value:
                    texts.append(value)
    except Exception:
        pass
    try:
        for series in chart.series:
            name = str(series.name).strip() if series.name is not None else ""
            if name:
                texts.append(name)
    except Exception:
        pass
    return texts


def _text_present(expected: str, haystack: list[str]) -> bool:
    needle = expected.strip()
    if not needle:
        return True
    if needle in haystack:
        return True
    # 允许段落合并或单元格拼接后的包含关系
    return any(needle in item for item in haystack)


def _shape_label(shape: BaseShape) -> str:
    name = getattr(shape, "name", None) or "未命名形状"
    return str(name)


def _preview(text: str, limit: int = 40) -> str:
    compact = " ".join(text.split())
    if len(compact) <= limit:
        return compact
    return compact[: limit - 1] + "…"
