from io import BytesIO

import pytest
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.ingest.base import UnsupportedDocument
from app.ingest.docx import DocxParser
from app.ingest.markdown import MarkdownParser
from app.ingest.pdf import PdfParser
from app.ingest.plain import PlainTextParser
from app.ingest.registry import SUPPORTED_EXTENSIONS, parser_for


def test_plain_text_splits_blank_line_paragraphs() -> None:
    data = "第一段内容。\n\n第二段内容。\n仍属第二段。\n\n第三段。".encode()
    parsed = PlainTextParser().parse(data)

    assert len(parsed.sections) == 3
    assert all(section.level == 0 for section in parsed.sections)
    assert parsed.sections[0].text == "第一段内容。"
    assert parsed.sections[1].text == "第二段内容。\n仍属第二段。"
    assert parsed.sections[2].locator == "第 3 段"


def test_plain_text_gbk_fallback() -> None:
    data = "中文纯文本".encode("gbk")
    parsed = PlainTextParser().parse(data)
    assert parsed.sections[0].text == "中文纯文本"
    assert parsed.warnings == []


def test_markdown_preserves_heading_levels_and_cleans_markup() -> None:
    source = """\
前言段落，含 **粗体** 与 `代码`。

# 一级标题

正文含 [链接文字](https://example.com) 与图片 ![示意图](a.png)。

## 二级标题

更多说明。

```
code_block_line
```
"""
    parsed = MarkdownParser().parse(source.encode())

    assert len(parsed.sections) == 3
    assert parsed.sections[0].level == 0
    assert parsed.sections[0].heading is None
    assert "粗体" in parsed.sections[0].text
    assert "**" not in parsed.sections[0].text
    assert "`" not in parsed.sections[0].text

    assert parsed.sections[1].level == 1
    assert parsed.sections[1].heading == "一级标题"
    assert "链接文字" in parsed.sections[1].text
    assert "https://example.com" not in parsed.sections[1].text
    assert "示意图" in parsed.sections[1].text
    assert "![" not in parsed.sections[1].text
    assert parsed.sections[1].locator.startswith("第 ")

    assert parsed.sections[2].level == 2
    assert parsed.sections[2].heading == "二级标题"
    assert "更多说明。" in parsed.sections[2].text
    assert "code_block_line" in parsed.sections[2].text


def test_markdown_setext_heading() -> None:
    source = "标题\n====\n\n正文\n"
    parsed = MarkdownParser().parse(source.encode())
    assert parsed.sections[0].level == 1
    assert parsed.sections[0].heading == "标题"
    assert parsed.sections[0].text == "正文"


def _build_docx() -> bytes:
    document = Document()
    document.add_heading("项目概述", level=1)
    document.add_paragraph("这是概述正文。")
    document.add_heading("关键指标", level=2)
    document.add_paragraph("指标说明。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "完成率"
    table.cell(1, 1).text = "95%"
    document.add_paragraph("")  # 空段落应被跳过
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_headings_and_table() -> None:
    parsed = DocxParser().parse(_build_docx())

    assert len(parsed.sections) == 2
    assert parsed.sections[0].level == 1
    assert parsed.sections[0].heading == "项目概述"
    assert "这是概述正文。" in parsed.sections[0].text
    assert parsed.sections[0].locator == "第 1 段"

    assert parsed.sections[1].level == 2
    assert parsed.sections[1].heading == "关键指标"
    assert "指标说明。" in parsed.sections[1].text
    assert "指标 | 数值" in parsed.sections[1].text
    assert "完成率 | 95%" in parsed.sections[1].text


def _build_pdf_with_headings() -> bytes:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    pdf.setFont("Helvetica-Bold", 24)
    pdf.drawString(72, height - 72, "Quarterly Review")
    pdf.setFont("Helvetica", 12)
    pdf.drawString(72, height - 110, "Body text on page one explains the summary.")
    pdf.drawString(72, height - 128, "Another body line stays at the normal size.")
    pdf.showPage()

    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(72, height - 72, "Details")
    pdf.setFont("Helvetica", 12)
    pdf.drawString(72, height - 110, "Page two continues with supporting details.")
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _build_empty_text_pdf() -> bytes:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def test_pdf_heading_heuristic_and_locators() -> None:
    parsed = PdfParser().parse(_build_pdf_with_headings())

    assert len(parsed.sections) >= 2
    headings = [section.heading for section in parsed.sections if section.heading]
    assert "Quarterly Review" in headings
    assert "Details" in headings

    review = next(section for section in parsed.sections if section.heading == "Quarterly Review")
    details = next(section for section in parsed.sections if section.heading == "Details")
    assert review.level == 1
    assert details.level == 2
    assert review.locator == "第 1 页"
    assert details.locator == "第 2 页"
    assert "Body text on page one" in review.text
    assert "Page two continues" in details.text


def test_pdf_scanned_warning() -> None:
    parsed = PdfParser().parse(_build_empty_text_pdf())
    assert parsed.sections == []
    assert parsed.warnings
    assert "扫描" in parsed.warnings[0]
    assert "OCR" in parsed.warnings[0]


def test_parser_for_by_extension() -> None:
    assert isinstance(parser_for("notes.TXT"), PlainTextParser)
    assert isinstance(parser_for("readme.markdown"), MarkdownParser)
    assert isinstance(parser_for("spec.docx"), DocxParser)
    assert isinstance(parser_for("deck.PDF"), PdfParser)
    assert ".md" in SUPPORTED_EXTENSIONS
    assert ".pdf" in SUPPORTED_EXTENSIONS


def test_unsupported_extension() -> None:
    with pytest.raises(UnsupportedDocument, match="不支持的文件类型"):
        parser_for("archive.zip")
