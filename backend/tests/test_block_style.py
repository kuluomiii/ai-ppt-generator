import uuid
from collections.abc import AsyncGenerator

import pytest
from httpx import ASGITransport, AsyncClient
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.core.db import async_session_factory
from app.domain.block_style import BlockStyle, merge_text_style, resolve_box
from app.domain.content import Deck, Slide, TextBlock
from app.domain.theme import get_theme
from app.domain.validation import validate_slide
from app.main import app
from app.models.project import Project, ProjectOutline
from app.models.slide import Slide as SlideRow
from app.render.pptx import render_deck_to_pptx


@pytest.fixture
async def client() -> AsyncGenerator[AsyncClient, None]:
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as value:
        yield value


async def _sign_up(client: AsyncClient) -> dict[str, str]:
    response = await client.post(
        "/api/v1/auth/register",
        json={"email": f"style_{uuid.uuid4().hex}@example.com", "password": "password123"},
    )
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


async def _project_with_slide(
    client: AsyncClient,
    headers: dict[str, str],
    *,
    blocks: list[dict] | None = None,
) -> tuple[dict, SlideRow]:
    response = await client.post(
        "/api/v1/projects",
        json={"title": "元素样式", "page_count": 5},
        headers=headers,
    )
    assert response.status_code == 201, response.text
    project = response.json()

    page_id = uuid.uuid4()
    slide_blocks = blocks or [
        {
            "id": "t1",
            "slot_id": "title",
            "type": "text",
            "text": "标题文字",
            "locked": False,
        },
        {
            "id": "b1",
            "slot_id": "body",
            "type": "bullets",
            "items": ["要点一", "要点二"],
            "locked": False,
        },
    ]

    async with async_session_factory() as session:
        result = await session.execute(
            select(Project)
            .options(selectinload(Project.outline))
            .where(Project.id == uuid.UUID(project["id"]))
        )
        record = result.scalar_one()
        session.add(
            ProjectOutline(
                project_id=record.id,
                status="confirmed",
                pages=[
                    {
                        "id": str(page_id),
                        "title": "页",
                        "objective": "目标",
                        "key_points": ["要点"],
                        "source_refs": [],
                        "layout_id": "bullets",
                    }
                ],
                revision=1,
            )
        )
        row = SlideRow(
            project_id=record.id,
            outline_page_id=page_id,
            position=1,
            layout_id="bullets",
            title="页",
            status="ready",
            blocks=slide_blocks,
            revision=1,
        )
        session.add(row)
        record.status = "ready"
        await session.commit()
        await session.refresh(row)
        return project, row


def test_merge_text_style_passthrough() -> None:
    theme = get_theme("ivory")
    merged = merge_text_style(theme, "title", None)
    base = theme.text_style("title")
    assert merged.size_pt == base.size_pt
    assert merged.color == base.color
    assert merged.italic is False


def test_merge_text_style_overrides() -> None:
    theme = get_theme("ivory")
    style = BlockStyle(
        size_pt=28,
        color="accent",
        weight=700,
        italic=True,
    )
    merged = merge_text_style(theme, "title", style)
    assert merged.size_pt == 28
    assert merged.color == "accent"
    assert merged.weight == 700
    assert merged.italic is True
    assert theme.color(merged.color) == theme.palette.accent


def test_merge_text_style_hex_color() -> None:
    theme = get_theme("ivory")
    style = BlockStyle(color="#0D9488")
    merged = merge_text_style(theme, "body", style)
    assert merged.color == "#0D9488"
    assert theme.color(merged.color) == "#0D9488"


def test_block_style_rejects_out_of_range() -> None:
    with pytest.raises(ValidationError):
        BlockStyle(size_pt=100)
    with pytest.raises(ValidationError):
        BlockStyle(color="not-a-color")
    with pytest.raises(ValidationError):
        BlockStyle(weight=100)


def test_resolve_box_fill_and_border() -> None:
    theme = get_theme("midnight")
    box = resolve_box(
        theme,
        BlockStyle(
            fill="accent_soft",
            radius_pt=8,
            border_color="accent",
            border_width_pt=1.5,
            padding_pt=12,
        ),
    )
    assert box.fill == theme.palette.accent_soft
    assert box.radius_pt == 8
    assert box.has_border
    assert box.border_color == theme.palette.accent
    assert box.padding_pt == 12


def test_overflow_uses_merged_size() -> None:
    from app.domain.content import BulletsBlock

    theme = get_theme("ivory")
    slide = Slide(
        id="s1",
        layout_id="bullets",
        blocks=[
            TextBlock(
                id="t1",
                slot_id="title",
                text="这是一段很长很长很长很长很长很长很长很长的标题内容",
                style=BlockStyle(size_pt=72),
            ),
            BulletsBlock(id="b1", slot_id="body", items=["要点"]),
        ],
    )
    issues = validate_slide(slide, theme=theme)
    assert any(issue.severity == "warning" and "溢出" in issue.message for issue in issues)


def test_pptx_applies_block_style() -> None:
    theme = get_theme("ivory")
    deck = Deck(
        id="d1",
        title="样式导出",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="cover",
                blocks=[
                    TextBlock(
                        id="t1",
                        slot_id="title",
                        text="加粗居中",
                        style=BlockStyle(
                            size_pt=40,
                            weight=700,
                            italic=True,
                            align="center",
                            color="#112233",
                        ),
                    ),
                    TextBlock(
                        id="s1",
                        slot_id="subtitle",
                        text="副标题",
                    ),
                ],
            )
        ],
    )
    buffer = render_deck_to_pptx(deck, theme=theme)
    presentation = Presentation(buffer)
    slide = presentation.slides[0]

    # 找到含「加粗居中」的段落
    hit = None
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs)
            if "加粗居中" in text:
                hit = paragraph
                break
        if hit is not None:
            break

    assert hit is not None
    assert hit.alignment == PP_ALIGN.CENTER
    run = hit.runs[0]
    assert run.font.size == Pt(40)
    assert run.font.bold is True
    assert run.font.italic is True
    from app.render.color import to_rgb

    assert run.font.color.rgb == to_rgb("#112233")


@pytest.mark.asyncio
async def test_style_api_does_not_lock(client: AsyncClient) -> None:
    headers = await _sign_up(client)
    project, slide = await _project_with_slide(client, headers)

    response = await client.patch(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/t1/style",
        headers=headers,
        json={
            "revision": slide.revision,
            "style": {"size_pt": 36, "color": "accent", "weight": 700},
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["revision"] == slide.revision + 1
    text_block = next(block for block in body["blocks"] if block["id"] == "t1")
    assert text_block["locked"] is False
    assert text_block["style"]["size_pt"] == 36
    assert text_block["style"]["color"] == "accent"


@pytest.mark.asyncio
async def test_style_api_clear_and_conflict(client: AsyncClient) -> None:
    headers = await _sign_up(client)
    project, slide = await _project_with_slide(
        client,
        headers,
        blocks=[
            {
                "id": "t1",
                "slot_id": "title",
                "type": "text",
                "text": "标题",
                "locked": False,
                "style": {"size_pt": 40, "italic": True},
            },
            {
                "id": "b1",
                "slot_id": "body",
                "type": "bullets",
                "items": ["要点"],
                "locked": False,
            },
        ],
    )

    cleared = await client.patch(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/t1/style",
        headers=headers,
        json={"revision": slide.revision, "style": None},
    )
    assert cleared.status_code == 200, cleared.text
    text_block = next(block for block in cleared.json()["blocks"] if block["id"] == "t1")
    assert text_block["style"] is None

    conflict = await client.patch(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/t1/style",
        headers=headers,
        json={"revision": slide.revision, "style": {"size_pt": 20}},
    )
    assert conflict.status_code == 409


@pytest.mark.asyncio
async def test_style_api_rejects_chart_and_invalid(client: AsyncClient) -> None:
    headers = await _sign_up(client)
    project, slide = await _project_with_slide(
        client,
        headers,
        blocks=[
            {
                "id": "t1",
                "slot_id": "title",
                "type": "text",
                "text": "图",
                "locked": False,
            },
            {
                "id": "c1",
                "slot_id": "chart",
                "type": "chart",
                "chart_type": "bar",
                "categories": ["A"],
                "series": [{"name": "S", "values": [1]}],
                "locked": False,
            },
        ],
    )
    # chart 布局才有 chart 槽；这里用 bullets 布局会让校验出问题，换 layout
    # 直接测 API 对 chart 类型的拒绝：先把 slide 的 layout 改成 chart
    async with async_session_factory() as session:
        row = await session.get(SlideRow, slide.id)
        assert row is not None
        row.layout_id = "chart"
        row.blocks = [
            {
                "id": "t1",
                "slot_id": "title",
                "type": "text",
                "text": "图",
                "locked": False,
            },
            {
                "id": "c1",
                "slot_id": "chart",
                "type": "chart",
                "chart_type": "bar",
                "categories": ["A"],
                "series": [{"name": "S", "values": [1]}],
                "locked": False,
            },
        ]
        await session.commit()
        await session.refresh(row)
        revision = row.revision

    rejected = await client.patch(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/c1/style",
        headers=headers,
        json={"revision": revision, "style": {"size_pt": 20}},
    )
    assert rejected.status_code == 422

    invalid = await client.patch(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/t1/style",
        headers=headers,
        json={"revision": revision, "style": {"size_pt": 999}},
    )
    assert invalid.status_code == 422


@pytest.mark.asyncio
async def test_style_api_large_size_triggers_overflow_warning(client: AsyncClient) -> None:
    headers = await _sign_up(client)
    long_title = "这是一段明显偏长的标题文字内容用来触发溢出检测ABCDEF"
    project, slide = await _project_with_slide(
        client,
        headers,
        blocks=[
            {
                "id": "t1",
                "slot_id": "title",
                "type": "text",
                "text": long_title,
                "locked": False,
            },
            {
                "id": "b1",
                "slot_id": "body",
                "type": "bullets",
                "items": ["要点"],
                "locked": False,
            },
        ],
    )

    response = await client.patch(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/t1/style",
        headers=headers,
        json={"revision": slide.revision, "style": {"size_pt": 72}},
    )
    assert response.status_code == 200, response.text
    assert any(
        issue["severity"] == "warning" and "溢出" in issue["message"]
        for issue in response.json()["issues"]
    )
