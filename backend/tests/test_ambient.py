from io import BytesIO

import pytest
from pptx import Presentation
from pptx.util import Pt

from app.domain.ambient import (
    AMBIENT_SHAPE_PREFIX,
    CornerBracket,
    EdgeBand,
    Glow,
    HairlineGrid,
    Watermark,
    iter_ambient_shapes,
    scope_of,
)
from app.domain.color import mix_hex
from app.domain.geometry import CANVAS_WIDTH_PT, Rect
from app.domain.sample import load_sample_deck
from app.domain.slide_geometry import placed_by_block_id
from app.domain.theme import Theme, get_theme, load_themes
from app.render.pptx import PptxRenderer
from app.render.verify import verify_pptx


def _theme_with(*motifs) -> Theme:
    return get_theme("midnight").model_copy(update={"ambient": list(motifs)})


def test_scope_maps_cover_and_section_apart_from_content() -> None:
    assert scope_of("cover") == "cover"
    assert scope_of("section") == "section"
    assert scope_of("bullets") == "content"
    assert scope_of("未知布局") == "content"


def test_theme_without_ambient_emits_nothing() -> None:
    theme = get_theme("midnight").model_copy(update={"ambient": []})
    assert iter_ambient_shapes(theme, "bullets") == []


def test_scope_filters_motifs() -> None:
    theme = _theme_with(EdgeBand(motif="edge_band", scope=["cover"]))
    assert len(iter_ambient_shapes(theme, "cover")) == 1
    assert iter_ambient_shapes(theme, "bullets") == []


def test_strength_mixes_token_into_background() -> None:
    theme = _theme_with(EdgeBand(motif="edge_band", color="accent", strength=0.4))
    shape = iter_ambient_shapes(theme, "bullets")[0]
    assert shape.color == mix_hex(theme.palette.accent, theme.palette.background, 0.4)


def test_hairline_grid_draws_inner_lines_only() -> None:
    theme = _theme_with(HairlineGrid(motif="hairline_grid", columns=4, rows=3))
    shapes = iter_ambient_shapes(theme, "bullets")
    # 4 等分只有 3 条竖线，3 等分只有 2 条横线；边界不画
    assert len(shapes) == 3 + 2
    assert all(shape.kind == "rect" for shape in shapes)
    assert all(0 < shape.rect.x < 1 for shape in shapes[:3])


def test_corner_bracket_hugs_declared_corner() -> None:
    theme = _theme_with(
        CornerBracket(motif="corner_bracket", corner="bottom_right", size_pt=80, inset_pt=20)
    )
    horizontal, vertical = iter_ambient_shapes(theme, "bullets")
    assert horizontal.rect.bottom == pytest.approx(vertical.rect.bottom, abs=1e-6)
    assert horizontal.rect.right == pytest.approx(vertical.rect.right, abs=1e-6)
    assert horizontal.rect.w > vertical.rect.w
    assert vertical.rect.h > horizontal.rect.h


def test_glow_layers_deepen_toward_center() -> None:
    theme = _theme_with(Glow(motif="glow", cx=0.5, cy=0.5, radius_pt=120, layers=3, strength=0.6))
    shapes = iter_ambient_shapes(theme, "bullets")
    assert [shape.kind for shape in shapes] == ["ellipse"] * 3
    # 越靠中心的层越小、颜色越浓，因此三层各不相同
    assert [shape.rect.w for shape in shapes] == sorted(
        (shape.rect.w for shape in shapes), reverse=True
    )
    assert len({shape.color for shape in shapes}) == 3


def test_watermark_falls_back_to_page_number() -> None:
    rect = {"x": 0.7, "y": 0.6, "w": 0.2, "h": 0.3}
    numbered = _theme_with(Watermark(motif="watermark", rect=rect))
    assert iter_ambient_shapes(numbered, "bullets", 0)[0].text == "01"
    assert iter_ambient_shapes(numbered, "bullets", 11)[0].text == "12"

    fixed = _theme_with(Watermark(motif="watermark", rect=rect, text="ACME"))
    assert iter_ambient_shapes(fixed, "bullets", 5)[0].text == "ACME"


def test_rect_motifs_are_clipped_into_canvas() -> None:
    """轴对齐矩形按外接框收边即是精确裁剪，几何无损，因此生成时就做掉。"""
    theme = _theme_with(
        # 臂长远超画布：两条臂都会被收边，收完仍是同样的矩形
        CornerBracket(motif="corner_bracket", corner="bottom_right", size_pt=1200, inset_pt=0),
        HairlineGrid(motif="hairline_grid", columns=3, rows=3),
    )
    for shape in iter_ambient_shapes(theme, "bullets"):
        assert shape.rect.x >= 0 and shape.rect.y >= 0
        assert shape.rect.right <= 1 + 1e-9
        assert shape.rect.bottom <= 1 + 1e-9


def test_bleeding_glow_keeps_layers_concentric() -> None:
    """光晕按外接框收边会被压扁成偏心月牙，所以它宁可出血也不收边。"""
    theme = _theme_with(Glow(motif="glow", cx=1.0, cy=0.0, radius_pt=400, layers=3))
    shapes = iter_ambient_shapes(theme, "bullets")
    assert len(shapes) == 3
    for shape in shapes:
        assert shape.rect.x + shape.rect.w / 2 == pytest.approx(1.0, abs=1e-9)
        assert shape.rect.y + shape.rect.h / 2 == pytest.approx(0.0, abs=1e-9)
    # 真出血才有意义：至少一层要越出画布
    assert any(shape.rect.x < 0 or shape.rect.right > 1 for shape in shapes)


def test_offscreen_motif_is_dropped() -> None:
    theme = _theme_with(Glow(motif="glow", cx=-2.0, cy=-2.0, radius_pt=10, layers=1))
    assert iter_ambient_shapes(theme, "bullets") == []


def test_avoid_content_keeps_only_shapes_in_the_gutter() -> None:
    """网格逐条判定，压在正文上的线整条丢掉，留在栏间空隙的线保留。"""
    theme = _theme_with(HairlineGrid(motif="hairline_grid", columns=4, avoid_content=True))
    # 左右两栏各占 45%，中间 0.45–0.55 是空隙：4 等分的三条线里只有 x=0.5 落在缝里
    occupied = [
        Rect(x=0.0, y=0.0, w=0.45, h=1.0),
        Rect(x=0.55, y=0.0, w=0.45, h=1.0),
    ]
    kept = iter_ambient_shapes(theme, "bullets", 0, occupied)
    assert [pytest.approx(shape.rect.x + shape.rect.w / 2, abs=1e-6) for shape in kept] == [0.5]


def test_avoid_content_off_draws_under_content() -> None:
    theme = _theme_with(HairlineGrid(motif="hairline_grid", columns=4))
    occupied = [Rect(x=0.0, y=0.0, w=1.0, h=1.0)]
    assert len(iter_ambient_shapes(theme, "bullets", 0, occupied)) == 3


def test_avoid_content_without_occupied_keeps_everything() -> None:
    """拿不到几何的场景（缩略图、纯主题预览）不该把装饰全丢掉。"""
    theme = _theme_with(HairlineGrid(motif="hairline_grid", columns=4, avoid_content=True))
    assert len(iter_ambient_shapes(theme, "bullets")) == 3


def test_avoid_content_drops_watermark_colliding_with_content() -> None:
    rect = {"x": 0.75, "y": 0.8, "w": 0.24, "h": 0.19}
    theme = _theme_with(Watermark(motif="watermark", rect=rect, avoid_content=True))
    free = [Rect(x=0.05, y=0.08, w=0.6, h=0.6)]
    assert len(iter_ambient_shapes(theme, "bullets", 0, free)) == 1
    taken = [Rect(x=0.6, y=0.7, w=0.36, h=0.25)]
    assert iter_ambient_shapes(theme, "bullets", 0, taken) == []


def test_every_preset_theme_keeps_flat_shapes_in_canvas() -> None:
    """只有椭圆允许出血：矩形和文字越界要么是配置写错，要么是几何算错。"""
    for theme in load_themes().values():
        for layout_id in ("cover", "section", "bullets", "kpi"):
            for shape in iter_ambient_shapes(theme, layout_id, 3):
                if shape.kind == "ellipse":
                    continue
                assert shape.rect.right <= 1 + 1e-9, f"{theme.id}/{layout_id}"
                assert shape.rect.bottom <= 1 + 1e-9, f"{theme.id}/{layout_id}"


def test_export_verify_allows_bleeding_ambient() -> None:
    """装饰出血是有意行为，导出校验不该把它判成越界。"""
    deck = load_sample_deck().model_copy(update={"theme_id": "obsidian"})
    theme = get_theme("obsidian")
    payload = PptxRenderer(theme).render(deck).getvalue()

    presentation = Presentation(BytesIO(payload))
    bleeding = [
        shape
        for shape in presentation.slides[0].shapes
        if shape.name.startswith(AMBIENT_SHAPE_PREFIX)
        and (int(shape.left) < 0 or int(shape.left) + int(shape.width) > Pt(CANVAS_WIDTH_PT))
    ]
    assert bleeding, "obsidian 的光晕本来就该越出画布"

    report = verify_pptx(payload, deck)
    assert report.passed is True, [issue.message for issue in report.issues]


def test_export_draws_ambient_below_content() -> None:
    deck = load_sample_deck()
    theme = get_theme("obsidian")
    presentation = Presentation(PptxRenderer(theme).render(deck))

    slide = presentation.slides[0]
    occupied = [placed.rect for placed in placed_by_block_id(deck.slides[0]).values()]
    expected = iter_ambient_shapes(theme, deck.slides[0].layout_id, 0, occupied)
    assert expected, "obsidian 应该在封面上有氛围层"

    shapes = list(slide.shapes)
    # 背景铺底之后紧跟氛围层，再往后才是装饰与内容
    ambient = shapes[1 : 1 + len(expected)]
    for shape, source in zip(ambient, expected, strict=True):
        assert shape.left == pytest.approx(source.rect.to_emu()[0], abs=2)
        assert shape.width == pytest.approx(source.rect.to_emu()[2], abs=2)
