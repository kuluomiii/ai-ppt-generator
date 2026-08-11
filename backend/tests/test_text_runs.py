"""emoji 必须单独成 run 并换字体，否则 PowerPoint 里是方框。"""

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.domain.block_style import merge_text_style
from app.domain.theme import get_theme
from app.render.text import write_paragraph

THEME = get_theme("ivory")
STYLE = merge_text_style(THEME, "body", None)


def _paragraph():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Pt(200), Pt(100))
    return box.text_frame.paragraphs[0]


def _typefaces(run) -> dict[str, str | None]:
    rpr = run.font._rPr
    return {
        tag: element.get("typeface")
        for tag in ("a:latin", "a:ea", "a:cs")
        if (element := rpr.find(qn(tag))) is not None
    }


def test_plain_text_stays_one_run() -> None:
    paragraph = _paragraph()
    write_paragraph(paragraph, "增长复盘", THEME, STYLE)
    assert [run.text for run in paragraph.runs] == ["增长复盘"]
    assert _typefaces(paragraph.runs[0])["a:ea"] == THEME.fonts.body.pptx_east_asian


def test_emoji_is_split_into_its_own_run() -> None:
    paragraph = _paragraph()
    write_paragraph(paragraph, "🧠 认知升级 📈", THEME, STYLE)

    assert [run.text for run in paragraph.runs] == ["🧠", " 认知升级 ", "📈"]
    emoji_font = THEME.fonts.emoji
    for index in (0, 2):
        faces = _typefaces(paragraph.runs[index])
        assert faces["a:latin"] == emoji_font.pptx_latin
        assert faces["a:ea"] == emoji_font.pptx_east_asian
        # 有播放器按复杂文本脚本挑字体，cs 不覆盖会退回正文字体
        assert faces["a:cs"] == emoji_font.pptx_latin
    assert _typefaces(paragraph.runs[1])["a:ea"] == THEME.fonts.body.pptx_east_asian


def test_variation_selector_stays_with_its_emoji() -> None:
    """⚠️ 是「符号 + 变体选择符」两个码位，切开就会退化成黑白轮廓。"""
    paragraph = _paragraph()
    write_paragraph(paragraph, "⚠️注意", THEME, STYLE)
    assert [run.text for run in paragraph.runs] == ["⚠️", "注意"]


def test_emoji_run_keeps_paragraph_text_style() -> None:
    paragraph = _paragraph()
    write_paragraph(paragraph, "📈 增长", THEME, STYLE)
    for run in paragraph.runs:
        assert run.font.size == Pt(STYLE.size_pt)


def test_empty_text_still_produces_a_run() -> None:
    paragraph = _paragraph()
    run = write_paragraph(paragraph, "", THEME, STYLE)
    assert run.text == ""
