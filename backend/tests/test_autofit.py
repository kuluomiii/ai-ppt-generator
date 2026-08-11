"""导出侧的溢出兜底：normAutofit 的比例必须由我们算好写死。"""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

from app.domain.content import BulletsBlock, Deck, Slide, TextBlock
from app.render.pptx import render_deck_to_pptx
from app.render.text import MIN_FONT_SCALE, disable_autofit, shrink_text_to_fit


def _frame(width_pt: float = 200.0, height_pt: float = 100.0):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Pt(width_pt), Pt(height_pt))
    return box.text_frame


def _autofit(frame) -> dict[str, str]:
    element = frame._bodyPr.normAutofit
    assert element is not None, "没有写入 normAutofit"
    return dict(element.attrib)


def test_new_textbox_no_longer_grows_itself() -> None:
    """python-pptx 新建文本框自带 spAutoFit，会把框撑出我们算好的槽位。"""
    frame = _frame()
    assert frame._bodyPr.spAutoFit is not None
    disable_autofit(frame)
    assert frame._bodyPr.spAutoFit is None
    assert frame._bodyPr.noAutofit is not None


def test_content_that_fits_is_left_alone() -> None:
    frame = _frame()
    assert shrink_text_to_fit(frame, needed_pt=80, available_pt=100) is None
    assert frame._bodyPr.normAutofit is None


def test_slight_overflow_only_tightens_line_spacing() -> None:
    """行距少 10% 几乎看不出来，字号变小一眼就能看见，所以先动行距。"""
    frame = _frame()
    assert shrink_text_to_fit(frame, needed_pt=105, available_pt=100) == 1.0
    attrs = _autofit(frame)
    assert attrs["fontScale"] == "100000"
    assert 0 < int(attrs["lnSpcReduction"]) <= 10000


def test_larger_overflow_shrinks_font_after_line_spacing() -> None:
    frame = _frame()
    scale = shrink_text_to_fit(frame, needed_pt=140, available_pt=100)
    assert scale is not None
    # 行距先吃掉 10%，剩下的缺口交给字号
    assert scale == pytest.approx((100 / 140) / 0.9)
    attrs = _autofit(frame)
    assert attrs["lnSpcReduction"] == "10000"
    assert int(attrs["fontScale"]) < 100000


def test_font_scale_stops_at_the_floor() -> None:
    """缩过头就不是排版微调而是内容超载，那该由生成期的溢出告警去管。"""
    frame = _frame()
    assert shrink_text_to_fit(frame, needed_pt=1000, available_pt=100) == MIN_FONT_SCALE
    assert _autofit(frame)["fontScale"] == str(round(MIN_FONT_SCALE * 100000))


def test_export_writes_autofit_for_overlong_text() -> None:
    long_text = "复盘这一年的增长路径与组织变化，" * 12
    deck = Deck(
        id="d1",
        title="溢出",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="bullets",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text=long_text),
                    BulletsBlock(id="b1", slot_id="body", items=["要点"]),
                ],
            )
        ],
    )
    presentation = Presentation(render_deck_to_pptx(deck))

    scales = [
        int(shape.text_frame._bodyPr.normAutofit.get("fontScale", "100000"))
        for shape in presentation.slides[0].shapes
        if shape.has_text_frame and shape.text_frame._bodyPr.normAutofit is not None
    ]
    assert scales, "排不下的标题应该写入 normAutofit"
    assert min(scales) < 100000


def test_export_leaves_short_text_at_full_size() -> None:
    deck = Deck(
        id="d1",
        title="正常",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="bullets",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="增长复盘"),
                    BulletsBlock(id="b1", slot_id="body", items=["营收增长 37%"]),
                ],
            )
        ],
    )
    presentation = Presentation(render_deck_to_pptx(deck))

    for shape in presentation.slides[0].shapes:
        if not shape.has_text_frame:
            continue
        body_pr = shape.text_frame._bodyPr
        assert body_pr.spAutoFit is None, "文本框不该自己长高"
        assert body_pr.normAutofit is None, "装得下就不该缩字号"
