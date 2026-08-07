import pytest
from httpx import ASGITransport, AsyncClient
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Pt

from app.domain.content import (
    BulletsBlock,
    ChartBlock,
    ChartSeries,
    Deck,
    ImageBlock,
    Slide,
    TextBlock,
)
from app.domain.flex_layout import FlexContainer, FlexLeaf
from app.domain.geometry import CANVAS_HEIGHT_PT, CANVAS_WIDTH_PT
from app.domain.sample import load_sample_deck
from app.domain.theme import get_theme, load_themes
from app.domain.validation import has_blocking_issue, validate_slide
from app.main import app
from app.render.color import to_rgb
from app.render.pptx import render_deck_to_pptx
from app.render.table import NO_STYLE_NO_GRID

_CHART_TYPE_MAP = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}


@pytest.fixture
def presentation() -> Presentation:
    return Presentation(render_deck_to_pptx(load_sample_deck()))


@pytest.fixture
async def client() -> AsyncClient:
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac


def _all_shapes(presentation: Presentation):
    for slide in presentation.slides:
        yield from slide.shapes


def _all_runs(presentation: Presentation):
    for shape in _all_shapes(presentation):
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                yield from paragraph.runs


def test_canvas_matches_shared_geometry(presentation: Presentation) -> None:
    assert presentation.slide_width == Pt(CANVAS_WIDTH_PT)
    assert presentation.slide_height == Pt(CANVAS_HEIGHT_PT)


def test_every_slide_is_exported(presentation: Presentation) -> None:
    assert len(presentation.slides) == len(load_sample_deck().slides)


def test_export_contains_no_bitmap(presentation: Presentation) -> None:
    """可编辑性的核心断言：导出结果里不允许出现任何位图。

    一旦某页被渲染成图片，它在 PowerPoint 里就无法编辑，
    这是本项目要避免的首要失败模式。
    """
    pictures = [
        shape.shape_type
        for shape in _all_shapes(presentation)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert not pictures


def test_text_is_native_and_complete(presentation: Presentation) -> None:
    exported = {run.text for run in _all_runs(presentation)}
    deck = load_sample_deck()

    expected: set[str] = set()
    for slide in deck.slides:
        for block in slide.blocks:
            if block.type == "text":
                expected.add(block.text)
            elif block.type == "bullets":
                expected.update(block.items)
            elif block.type == "kpi":
                expected.update({block.value, block.label})

    assert expected <= exported, expected - exported


def test_chinese_font_is_declared(presentation: Presentation) -> None:
    """只设 a:latin 时中文会落到 PowerPoint 的兜底字体，导出效果与 Web 端不一致"""
    runs = list(_all_runs(presentation))
    assert runs

    for run in runs:
        east_asian = run.font._rPr.find(qn("a:ea"))
        assert east_asian is not None, f"未声明中日韩字体：{run.text}"
        assert east_asian.get("typeface")


def test_bullets_use_native_markers(presentation: Presentation) -> None:
    markers = [
        paragraph._p.find(qn("a:pPr"))
        for shape in _all_shapes(presentation)
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
    ]
    native = [
        ppr
        for ppr in markers
        if ppr is not None
        and (ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buAutoNum")) is not None)
    ]
    assert native, "要点应使用 PowerPoint 原生项目符号，而不是把符号拼进文本"


def test_table_is_a_native_table(presentation: Presentation) -> None:
    tables = [shape.table for shape in _all_shapes(presentation) if shape.has_table]
    assert len(tables) == 1

    table = tables[0]
    source = next(
        block
        for slide in load_sample_deck().slides
        for block in slide.blocks
        if block.type == "table"
    )
    assert len(table.columns) == len(source.header)
    assert len(table.rows) == len(source.rows) + 1
    assert table.cell(0, 0).text == source.header[0]


def test_table_uses_our_borders_not_powerpoint_defaults(presentation: Presentation) -> None:
    """不显式声明边框时，PowerPoint 会套用默认表格样式画出竖向网格线与斑马纹，
    导出观感与 Web 端不一致"""
    table = next(shape.table for shape in _all_shapes(presentation) if shape.has_table)

    style_id = table._tbl.find(qn("a:tblPr")).find(qn("a:tableStyleId"))
    assert style_id is not None and style_id.text == NO_STYLE_NO_GRID

    tc_pr = table.cell(1, 0)._tc.find(qn("a:tcPr"))
    assert tc_pr.find(qn("a:lnL")).find(qn("a:noFill")) is not None
    assert tc_pr.find(qn("a:lnR")).find(qn("a:noFill")) is not None
    assert tc_pr.find(qn("a:lnB")).find(qn("a:solidFill")) is not None


def test_image_placeholder_is_made_of_shapes(presentation: Presentation) -> None:
    """占位图必须由形状拼出。一旦生成位图，这块区域在 PowerPoint 里就不可编辑"""
    freeforms = [
        shape for shape in _all_shapes(presentation) if shape.shape_type == MSO_SHAPE_TYPE.FREEFORM
    ]
    image_blocks = [
        block
        for slide in load_sample_deck().slides
        for block in slide.blocks
        if block.type == "image"
    ]
    assert len(freeforms) == len(image_blocks)


@pytest.mark.parametrize("theme_id", sorted(load_themes()))
def test_every_theme_renders(theme_id: str) -> None:
    exported = Presentation(render_deck_to_pptx(load_sample_deck(), theme_id))
    assert list(exported.slides)


@pytest.mark.asyncio
async def test_download_endpoint_returns_pptx(client: AsyncClient) -> None:
    response = await client.get("/api/v1/design/sample-deck/pptx?theme_id=midnight")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    # PPTX 是 zip 容器，前两个字节固定为 PK
    assert response.content[:2] == b"PK"


@pytest.mark.asyncio
async def test_download_unknown_theme_returns_404(client: AsyncClient) -> None:
    response = await client.get("/api/v1/design/sample-deck/pptx?theme_id=nope")
    assert response.status_code == 404


def _chart_deck(
    chart_type: str,
    *,
    categories: list[str] | None = None,
    series: list[ChartSeries] | None = None,
    unit: str | None = "单位",
) -> Deck:
    return Deck(
        id="chart-test",
        title="图表测试",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="chart",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="图表页"),
                    ChartBlock(
                        id="c1",
                        slot_id="chart",
                        chart_type=chart_type,  # type: ignore[arg-type]
                        categories=categories or ["甲", "乙", "丙"],
                        series=series
                        or [
                            ChartSeries(name="系列一", values=[10, 20, 30]),
                            ChartSeries(name="系列二", values=[15, 25, 18]),
                        ],
                        unit=unit,
                    ),
                ],
            )
        ],
    )


def _first_chart(presentation: Presentation):
    for shape in presentation.slides[0].shapes:
        if shape.has_chart:
            return shape.chart
    raise AssertionError("未找到原生图表")


@pytest.mark.parametrize("chart_type", ["bar", "column", "line", "pie"])
def test_chart_exports_as_native_editable_object(chart_type: str) -> None:
    series = [ChartSeries(name="占比", values=[10, 20, 30])] if chart_type == "pie" else None
    presentation = Presentation(
        render_deck_to_pptx(_chart_deck(chart_type, series=series, unit=None))
    )
    chart = _first_chart(presentation)
    assert chart.chart_type == _CHART_TYPE_MAP[chart_type]


def test_chart_writes_categories_and_series() -> None:
    presentation = Presentation(render_deck_to_pptx(_chart_deck("column")))
    chart = _first_chart(presentation)

    assert list(chart.plots[0].categories) == ["甲", "乙", "丙"]
    assert [series.name for series in chart.series] == ["系列一", "系列二"]
    assert list(chart.series[0].values) == [10.0, 20.0, 30.0]
    assert list(chart.series[1].values) == [15.0, 25.0, 18.0]


def test_chart_series_color_comes_from_theme() -> None:
    theme = get_theme("ivory")
    presentation = Presentation(render_deck_to_pptx(_chart_deck("column"), "ivory"))
    chart = _first_chart(presentation)

    assert chart.series[0].format.fill.fore_color.rgb == to_rgb(theme.palette.chart_series[0])


def test_chart_hides_legend_for_single_series() -> None:
    deck = _chart_deck(
        "column",
        series=[ChartSeries(name="唯一", values=[1, 2, 3])],
    )
    chart = _first_chart(Presentation(render_deck_to_pptx(deck)))
    assert chart.has_legend is False


def test_chart_shows_legend_for_multiple_series() -> None:
    chart = _first_chart(Presentation(render_deck_to_pptx(_chart_deck("column"))))
    assert chart.has_legend is True


def test_pie_chart_colors_by_point() -> None:
    theme = get_theme("ivory")
    deck = _chart_deck(
        "pie",
        series=[ChartSeries(name="占比", values=[10, 20, 30])],
        unit=None,
    )
    chart = _first_chart(Presentation(render_deck_to_pptx(deck, "ivory")))
    colors = theme.palette.chart_series
    for index, point in enumerate(chart.series[0].points):
        assert point.format.fill.fore_color.rgb == to_rgb(colors[index % len(colors)])


def test_chart_mismatched_lengths_still_export() -> None:
    deck = _chart_deck(
        "bar",
        categories=["A", "B", "C", "D"],
        series=[ChartSeries(name="短", values=[1, 2])],
    )
    buffer = render_deck_to_pptx(deck)
    assert buffer.getvalue()[:2] == b"PK"

    chart = _first_chart(Presentation(buffer))
    assert list(chart.plots[0].categories) == ["A", "B"]
    assert list(chart.series[0].values) == [1.0, 2.0]


def test_chart_capacity_overflow_is_warning_not_error() -> None:
    slide = Slide(
        id="s1",
        layout_id="chart",
        blocks=[
            TextBlock(id="t1", slot_id="title", text="标题"),
            ChartBlock(
                id="c1",
                slot_id="chart",
                chart_type="column",
                categories=[f"C{i}" for i in range(10)],
                series=[ChartSeries(name=f"S{i}", values=[float(i)] * 10) for i in range(4)],
            ),
        ],
    )
    issues = validate_slide(slide)
    warnings = [issue for issue in issues if issue.severity == "warning"]
    errors = [issue for issue in issues if issue.severity == "error"]

    assert not errors
    assert any("系列" in issue.message for issue in warnings)
    assert any("分类" in issue.message for issue in warnings)


def _flex_image_left_slide() -> Slide:
    return Slide(
        id="flex-1",
        layout_id="image-left",
        layout_mode="flex",
        layout_tree=FlexContainer(
            type="row",
            id="root",
            gap_pt=16,
            ratios=[38, 62],
            children=[
                FlexLeaf(id="leaf-image", block_id="image"),
                FlexContainer(
                    type="column",
                    id="text-col",
                    gap_pt=16,
                    children=[
                        FlexLeaf(
                            id="leaf-title",
                            block_id="title",
                            grow=0.5,
                            text_style="title",
                        ),
                        FlexLeaf(
                            id="leaf-body",
                            block_id="body",
                            grow=1.5,
                            text_style="bullet",
                        ),
                    ],
                ),
            ],
        ),
        blocks=[
            ImageBlock(id="image", slot_id="image", alt="配图", source="placeholder"),
            TextBlock(id="title", slot_id="title", text="灵活布局标题"),
            BulletsBlock(id="body", slot_id="body", items=["要点一", "要点二"]),
        ],
    )


def test_flex_slide_validates_and_renders_pptx() -> None:
    slide = _flex_image_left_slide()
    issues = validate_slide(slide, theme_id="ivory")
    assert not has_blocking_issue(issues)

    deck = Deck(id="d-flex", title="flex", theme_id="ivory", slides=[slide])
    buffer = render_deck_to_pptx(deck, theme_id="ivory")
    assert buffer.getvalue()[:2] == b"PK"
    presentation = Presentation(buffer)
    assert len(presentation.slides) == 1
    texts = []
    for shape in presentation.slides[0].shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                texts.append("".join(run.text for run in paragraph.runs))
    joined = "\n".join(texts)
    assert "灵活布局标题" in joined
    assert "要点一" in joined


def test_flex_slide_with_preset_renders_pptx() -> None:
    slide = _flex_image_left_slide()
    assert slide.layout_tree is not None
    slide.layout_tree.preset = "solid_boxes"
    deck = Deck(id="d-flex-skin", title="flex-skin", theme_id="ivory", slides=[slide])
    buffer = render_deck_to_pptx(deck, theme_id="ivory")
    assert buffer.getvalue()[:2] == b"PK"
    presentation = Presentation(buffer)
    # 背景 + 2 个 fill_box 皮肤 + 内容形状，至少多于无皮肤时
    assert len(presentation.slides[0].shapes) >= 4
